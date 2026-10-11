"""Create the Word report from its reviewable Markdown source."""
import re
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT

ROOT=Path(__file__).resolve().parent
doc=Document()
section=doc.sections[0]
section.page_width=Inches(8.27);section.page_height=Inches(11.69)
section.top_margin=section.bottom_margin=Inches(.7)
section.left_margin=section.right_margin=Inches(.75)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2']:
    style=doc.styles[name]
    style.font.name='Calibri';style.font.color.rgb=RGBColor(0,0,0)
    style.font.size=Pt(11 if name=='Normal' else 25 if name=='Title' else 15 if name=='Heading 1' else 12)
    style.paragraph_format.space_after=Pt(7)
    style.paragraph_format.line_spacing=1.1
doc.styles['Heading 1'].paragraph_format.space_before=Pt(15)
doc.styles['Normal'].paragraph_format.widow_control=True
for style in doc.styles:
    for border in list(style.element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
footer=section.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run('SQL Sentinel  |  ')
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)

def table(rows):
    t=doc.add_table(rows=1,cols=len(rows[0]));t.alignment=WD_TABLE_ALIGNMENT.CENTER
    t.autofit=False
    n=len(rows[0]);widths=[2.55]+[(6.77-2.55)/(n-1)]*(n-1) if n in (3,6) else [6.77/n]*n
    for c,w in zip(t.columns,widths):c.width=Inches(w)
    props=t._tbl.tblPr
    borders=OxmlElement('w:tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:
        e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
    props.append(borders)
    for i,values in enumerate(rows):
        cells=t.rows[0].cells if i==0 else t.add_row().cells
        for j,(cell,value) in enumerate(zip(cells,values)):
            cell.width=Inches(widths[j]);cell.text=value;cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcPr=cell._tc.get_or_add_tcPr()
            shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'DCE6F1' if i==0 else 'F5F7FA' if i%2==0 else 'FFFFFF');tcPr.append(shade)
            margins=OxmlElement('w:tcMar')
            for edge in ['top','bottom','left','right']:
                e=OxmlElement('w:'+edge);e.set(qn('w:w'),'90');e.set(qn('w:type'),'dxa');margins.append(e)
            tcPr.append(margins)
            for p in cell.paragraphs:
                p.paragraph_format.space_after=Pt(3);p.paragraph_format.line_spacing=1.05
                p.paragraph_format.keep_with_next=(i>=len(rows)-3 and i<len(rows)-1)
                for r in p.runs:r.font.size=Pt(9);r.bold=i==0
        tr=t.rows[i]._tr.get_or_add_trPr()
        if i==0:tr.append(OxmlElement('w:tblHeader'))
        tr.append(OxmlElement('w:cantSplit'))
    doc.add_paragraph('').paragraph_format.space_after=Pt(2)

lines=(ROOT/'Project-Technical-Report.md').read_text(encoding='utf-8-sig').splitlines()
i=0
while i<len(lines):
    line=lines[i].strip()
    if not line:i+=1;continue
    if line.startswith('|'):
        rows=[]
        while i<len(lines) and lines[i].strip().startswith('|'):
            cells=[x.strip() for x in lines[i].strip().strip('|').split('|')]
            if not all(re.fullmatch(r':?-+:?',x) for x in cells):rows.append(cells)
            i+=1
        table(rows);continue
    if line.startswith('# '):doc.add_paragraph(line[2:],style='Title')
    elif line.startswith('## '):doc.add_paragraph(line[3:],style='Heading 1')
    else:doc.add_paragraph(line)
    i+=1
doc.core_properties.title='SQL Sentinel SQL Injection Detection Project Technical Report'
doc.core_properties.subject='CNS Project 76 algorithms training application evaluation and deployment'
doc.core_properties.author='SQL Sentinel Project Team'
doc.save(ROOT/'Project-Technical-Report.docx')
print('Word report created')
